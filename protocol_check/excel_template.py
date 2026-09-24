"""Excel 接收 CAN 协议模板：读写与示例生成"""

from typing import Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from protocol_check.models import MessageTemplate, SignalTemplate, TemplateBook

SHEET_MSG = "报文"
SHEET_SIG = "信号"
SHEET_DOC = "说明"

MSG_HEADERS = ["报文ID", "报文名称", "期望DLC", "必须出现", "说明"]
SIG_HEADERS = [
    "报文ID",
    "信号名",
    "起始位",
    "长度",
    "字节序",
    "类型",
    "因子",
    "偏移",
    "单位",
    "说明",
]

_HEADER_FILL = PatternFill("solid", fgColor="2563EB")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def _norm_id(value) -> int:
    if value is None or value == "":
        raise ValueError("报文ID为空")
    if isinstance(value, int):
        return value
    s = str(value).strip().replace("0x", "").replace("0X", "").replace(" ", "")
    if s.endswith("x") or s.endswith("X"):
        s = s[:-1]
    return int(s, 16)


def _norm_bool(value, default: bool = True) -> bool:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    s = str(value).strip().lower()
    return s in ("1", "y", "yes", "true", "是", "必须", "√", "x", "required")


def _norm_float(value, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    try:
        return float(value)
    except Exception:
        return default


def _style_header(ws, headers: List[str]) -> None:
    ws.append(headers)
    for col, _ in enumerate(headers, 1):
        cell = ws.cell(1, col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"


def _autosize(ws, min_w: int = 10, max_w: int = 28) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = 0
        for cell in col[:80]:
            v = "" if cell.value is None else str(cell.value)
            length = max(length, len(v) + 2)
        ws.column_dimensions[letter].width = max(min_w, min(max_w, length))


def _find_sheet(wb, names: List[str]):
    lower_map = {name.lower(): name for name in wb.sheetnames}
    for n in names:
        if n in wb.sheetnames:
            return wb[n]
        if n.lower() in lower_map:
            return wb[lower_map[n.lower()]]
    return None


def _header_index(ws) -> Dict[str, int]:
    mapping = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(1, col).value
        if val is None:
            continue
        mapping[str(val).strip().lower()] = col
    return mapping


def load_template_excel(path: str) -> TemplateBook:
    wb = openpyxl.load_workbook(path, data_only=True)
    book = TemplateBook(source_path=path)

    ws_msg = _find_sheet(wb, [SHEET_MSG, "messages", "message", "报文定义"])
    ws_sig = _find_sheet(wb, [SHEET_SIG, "signals", "signal", "信号定义"])

    if ws_msg is None and ws_sig is None:
        raise ValueError("Excel 中未找到「报文」或「信号」工作表")

    if ws_msg is not None:
        idx = _header_index(ws_msg)
        col_id = idx.get("报文id", 1)
        col_name = idx.get("报文名称", idx.get("名称", 2))
        col_dlc = idx.get("期望dlc", idx.get("dlc", 3))
        col_req = idx.get("必须出现", idx.get("必须", 4))
        col_note = idx.get("说明", idx.get("备注", 5))
        for row in range(2, ws_msg.max_row + 1):
            raw_id = ws_msg.cell(row, col_id).value
            if raw_id is None or str(raw_id).strip() == "":
                continue
            try:
                msg_id = _norm_id(raw_id)
            except Exception:
                continue
            name = ws_msg.cell(row, col_name).value
            dlc_raw = ws_msg.cell(row, col_dlc).value
            try:
                dlc = int(float(dlc_raw)) if dlc_raw not in (None, "") else 8
            except Exception:
                dlc = 8
            book.add_message(
                MessageTemplate(
                    msg_id=msg_id,
                    name="" if name is None else str(name).strip(),
                    expected_dlc=dlc,
                    required=_norm_bool(ws_msg.cell(row, col_req).value, True),
                    note="" if ws_msg.cell(row, col_note).value is None else str(ws_msg.cell(row, col_note).value),
                )
            )

    if ws_sig is not None:
        idx = _header_index(ws_sig)
        col_id = idx.get("报文id", 1)
        col_name = idx.get("信号名", idx.get("名称", 2))
        col_start = idx.get("起始位", 3)
        col_len = idx.get("长度", 4)
        col_order = idx.get("字节序", 5)
        col_type = idx.get("类型", 6)
        col_factor = idx.get("因子", 7)
        col_offset = idx.get("偏移", 8)
        col_unit = idx.get("单位", 9)
        col_note = idx.get("说明", idx.get("备注", 10))
        for row in range(2, ws_sig.max_row + 1):
            raw_id = ws_sig.cell(row, col_id).value
            raw_name = ws_sig.cell(row, col_name).value
            if raw_id is None or str(raw_id).strip() == "":
                continue
            if raw_name is None or str(raw_name).strip() == "":
                continue
            try:
                msg_id = _norm_id(raw_id)
            except Exception:
                continue
            type_raw = ws_sig.cell(row, col_type).value
            type_s = "" if type_raw is None else str(type_raw).strip().lower()
            signed = type_s in ("int", "signed", "有符号", "sint")
            order_raw = ws_sig.cell(row, col_order).value
            order = "Intel" if order_raw in (None, "") else str(order_raw).strip()
            if order.lower() in ("le", "l", "小端", "intel"):
                order = "Intel"
            elif order.lower() in ("be", "m", "大端", "motorola"):
                order = "Motorola"
            try:
                start = int(float(ws_sig.cell(row, col_start).value))
            except Exception:
                start = 0
            try:
                length = int(float(ws_sig.cell(row, col_len).value))
            except Exception:
                length = 8
            book.add_signal(
                SignalTemplate(
                    msg_id=msg_id,
                    name=str(raw_name).strip(),
                    start_bit=start,
                    length=length,
                    byte_order=order,
                    signed=signed,
                    factor=_norm_float(ws_sig.cell(row, col_factor).value, 1.0),
                    offset=_norm_float(ws_sig.cell(row, col_offset).value, 0.0),
                    unit="" if ws_sig.cell(row, col_unit).value is None else str(ws_sig.cell(row, col_unit).value).strip(),
                    note="" if ws_sig.cell(row, col_note).value is None else str(ws_sig.cell(row, col_note).value),
                )
            )

    wb.close()
    return book


def save_template_excel(book: TemplateBook, path: str, with_doc: bool = True) -> None:
    wb = openpyxl.Workbook()
    # 说明
    if with_doc:
        ws_doc = wb.active
        ws_doc.title = SHEET_DOC
        doc_rows = [
            ["BMS 接收 CAN 协议核对模板"],
            [""],
            ["用途", "作为「协议核对」Tab 的判断标准：实报 ID/DLC/信号位段与此表对照"],
            ["报文ID", "十六进制，可写 18A2F2F2 或 0x18A2F2F2"],
            ["期望DLC", "数据字节数，通常为 8"],
            ["必须出现", "是：日志/采集中缺失则报异常；否：允许不出现"],
            ["信号表", "可为空；为空时只核 ID + DLC"],
            ["起始位", "Intel：信号 LSB 的 bit 序号；Motorola：信号 MSB 的 bit 序号（Vector DBC 习惯）"],
            ["字节序", "Intel（@1 小端）或 Motorola（@0 大端，按 DBC 锯齿位序）"],
            ["DBC导入", "编辑器「导入 DBC」默认把 0x98… 映射为总线 0x18…（&0x1FFFFFFF），可在表中改"],
            ["类型", "uint / int（int 表示有符号）"],
            ["核对规则", "1) 实报ID是否在模板 2) 必须出现的ID是否缺失 3) DLC是否一致 4) 信号位段是否可解码"],
        ]
        for r in doc_rows:
            ws_doc.append(r)
        ws_doc["A1"].font = Font(bold=True, size=14, color="2563EB")
        _autosize(ws_doc, min_w=12, max_w=70)

    ws_msg = wb.create_sheet(SHEET_MSG)
    _style_header(ws_msg, MSG_HEADERS)
    for msg in sorted(book.messages.values(), key=lambda m: m.msg_id):
        ws_msg.append(
            [
                f"{msg.msg_id:X}",
                msg.name,
                msg.expected_dlc,
                "是" if msg.required else "否",
                msg.note,
            ]
        )
    _autosize(ws_msg)

    ws_sig = wb.create_sheet(SHEET_SIG)
    _style_header(ws_sig, SIG_HEADERS)
    rows = []
    for msg in book.messages.values():
        for sig in msg.signals:
            rows.append(sig)
    rows.sort(key=lambda s: (s.msg_id, s.start_bit, s.name))
    for sig in rows:
        ws_sig.append(
            [
                f"{sig.msg_id:X}",
                sig.name,
                sig.start_bit,
                sig.length,
                sig.byte_order,
                "int" if sig.signed else "uint",
                sig.factor,
                sig.offset,
                sig.unit,
                sig.note,
            ]
        )
    _autosize(ws_sig)

    wb.save(path)
    wb.close()


def _ids_from_asc(asc_path: str) -> List[int]:
    """从 Vector ASC 抽取唯一 CAN ID（十六进制），按数值排序。"""
    import re

    pat = re.compile(r"^\s*[\d.]+\s+(?:\d+|---)\s+([0-9A-Fa-f]+)x?\s+(?:Rx|Tx)", re.IGNORECASE)
    ids = set()
    with open(asc_path, "r", encoding="utf-8", errors="ignore") as f:
        for line in f:
            if "Status:" in line:
                continue
            m = pat.match(line)
            if m:
                ids.add(int(m.group(1), 16))
    return sorted(ids)


def build_sample_template(asc_path: Optional[str] = None) -> TemplateBook:
    """生成示例接收模板。

    若提供 asc_path，则把日志中出现过的 ID 全部写入模板（便于立刻跑通核对）；
    信号表只给少数报文示例行，其余仅核 ID+DLC。
    """
    book = TemplateBook(source_path="sample")

    # 少数报文挂示例信号，方便看字段级核对
    example_signals = {
        0x18A2F2F2: [
            SignalTemplate(0x18A2F2F2, "Byte0_状态", 0, 8, "Intel", False, 1, 0, "", "示例"),
            SignalTemplate(0x18A2F2F2, "Byte1_预警", 8, 8, "Intel", False, 1, 0, "", "示例"),
        ],
        0x18A3FFA1: [
            SignalTemplate(0x18A3FFA1, "Byte0_状态", 0, 8, "Intel", False, 1, 0, "", "示例"),
        ],
        0x18B2FFA1: [
            SignalTemplate(0x18B2FFA1, "示例_u16", 0, 16, "Intel", False, 0.1, 0, "V", "示例"),
        ],
    }

    name_hints = {
        0x18A2F2F2: "Warning/环流相关(硬件常见ID)",
        0x18A3FFA1: "WarningInfo(协议标准ID)",
        0x3C2: "标准帧/心跳类",
    }

    ids: List[int] = []
    if asc_path:
        try:
            ids = _ids_from_asc(asc_path)
        except Exception:
            ids = []
    if not ids:
        ids = [
            0x18A2F2F2, 0x18A3FFA1, 0x18B0FFA1, 0x18B1FFA1, 0x18B2FFA1,
            0x18B5FFA1, 0x18B7FFA1, 0x18BBFFA1, 0x18A2FFB1, 0x18B8FFB1,
        ]

    for msg_id in ids:
        name = name_hints.get(msg_id, f"CAN_{msg_id:X}")
        book.add_message(
            MessageTemplate(
                msg_id=msg_id,
                name=name,
                expected_dlc=8,
                required=True,
                note="示例",
            )
        )
        for sig in example_signals.get(msg_id, []):
            book.add_signal(sig)

    return book
