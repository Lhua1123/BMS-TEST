"""
报告导出模块
支持Excel报告导出功能
"""

import logging
import os
from datetime import datetime
from typing import List, Dict, Any, Optional
from pathlib import Path

logger = logging.getLogger(__name__)


class ReportExporter:
    """报告导出器"""
    
    def __init__(self, template_path: Optional[str] = None):
        self.template_path = template_path
        
    def export_batch_report(self, report: Dict[str, Any], output_path: str) -> bool:
        """Export one batch report."""
        return self.export_batch_reports([report], output_path)

    def export_batch_reports(self, reports: List[Dict[str, Any]], output_path: str) -> bool:
        """Export batch summaries and all nested test details to one workbook."""
        try:
            import openpyxl
            from openpyxl.styles import Alignment, Font, PatternFill

            wb = openpyxl.Workbook()
            summary_ws = wb.active
            summary_ws.title = "\u62a5\u544a\u6c47\u603b"
            detail_ws = wb.create_sheet("\u6d4b\u8bd5\u660e\u7ec6")

            summary_headers = ["\u6d4b\u8bd5\u65f6\u95f4", "\u6d4b\u8bd5\u7c7b\u578b", "\u6d4b\u8bd5\u6570\u91cf", "\u901a\u8fc7", "\u5931\u8d25", "\u901a\u8fc7\u7387", "\u72b6\u6001", "\u8017\u65f6(\u79d2)"]
            detail_headers = ["\u62a5\u544a\u65f6\u95f4", "\u6d4b\u8bd5\u7c7b\u578b", "\u6d4b\u8bd5\u9879\u76ee", "\u7f16\u53f7", "\u72b6\u6001", "\u8017\u65f6(\u79d2)", "\u9519\u8bef\u4fe1\u606f", "\u68c0\u67e5\u9879\u8be6\u60c5"]
            for col, value in enumerate(summary_headers, 1):
                cell = summary_ws.cell(row=1, column=col, value=value)
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for col, value in enumerate(detail_headers, 1):
                cell = detail_ws.cell(row=1, column=col, value=value)
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")

            detail_row = 2
            for summary_row, report in enumerate(reports, 2):
                start_timestamp = report.get("start_timestamp")
                time_text = datetime.fromtimestamp(start_timestamp).strftime("%Y-%m-%d %H:%M:%S") if start_timestamp else "-"
                total = report.get("total_count", 0) or 0
                passed = report.get("passed_count", 0) or 0
                failed = report.get("failed_count", 0) or 0
                rate = passed / total * 100 if total else 0.0
                status_text = "通过" if report.get("status") in ("PASS", "PASSED", "通过") else "失败"
                test_type = {"fault": "\u6545\u969c\u7801\u6d4b\u8bd5", "threshold": "\u9608\u503c\u6d4b\u8bd5"}.get(
                    report.get("test_type"), report.get("test_type", "")
                )
                summary_values = [
                    time_text, test_type, total, passed, failed,
                    f"{rate:.1f}%", status_text, f"{report.get('duration_sec', 0):.2f}",
                ]
                for col, value in enumerate(summary_values, 1):
                    summary_ws.cell(row=summary_row, column=col, value=value)
                summary_ws.cell(row=summary_row, column=7).fill = PatternFill(
                    start_color="C6EFCE" if status_text == "\u901a\u8fc7" else "FFC7CE",
                    end_color="C6EFCE" if status_text == "\u901a\u8fc7" else "FFC7CE",
                    fill_type="solid",
                )

                for result in report.get("details", []) or []:
                    metadata = result.get("metadata") or {}
                    name = result.get("fault_name") or result.get("pid_name") or "-"
                    code = metadata.get("fault_code") or metadata.get("pid") or result.get("test_id") or "-"
                    raw_status = result.get("overall_status", "")
                    result_status = "通过" if raw_status in ("通过", "PASSED", "PASS", "成功") else "失败"
                    checks = []
                    for check in result.get("checks", []) or []:
                        checks.append(
                            f"{check.get('type') or check.get('name')}: "
                            f"\u671f\u671b={check.get('expected_value', '-')}, "
                            f"\u5b9e\u9645={check.get('actual_value', '-')}, "
                            f"\u72b6\u6001={check.get('status', '')}, "
                            f"{check.get('message') or ''}"
                        )
                    detail_values = [
                        time_text, test_type, name, code, result_status,
                        f"{result.get('duration', 0):.2f}",
                        result.get("error_message") or "",
                        "\n".join(checks),
                    ]
                    for col, value in enumerate(detail_values, 1):
                        cell = detail_ws.cell(row=detail_row, column=col, value=value)
                        cell.alignment = Alignment(vertical="top", wrap_text=True)
                    detail_ws.cell(row=detail_row, column=5).fill = PatternFill(
                        start_color="C6EFCE" if result_status == "\u901a\u8fc7" else "FFC7CE",
                        end_color="C6EFCE" if result_status == "\u901a\u8fc7" else "FFC7CE",
                        fill_type="solid",
                    )
                    detail_row += 1

            summary_widths = [22, 16, 12, 10, 10, 12, 10, 14]
            detail_widths = [22, 16, 28, 18, 10, 14, 36, 70]
            for index, width in enumerate(summary_widths, 1):
                summary_ws.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
            for index, width in enumerate(detail_widths, 1):
                detail_ws.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
            summary_ws.freeze_panes = "A2"
            detail_ws.freeze_panes = "A2"
            wb.save(output_path)
            return True
        except Exception as e:
            logger.error(f"Batch report export failed: {e}", exc_info=True)
            return False

    def export_to_excel(self, test_results: List[Dict[str, Any]], 
                       output_path: str, 
                       company_name: str = "BMS测试团队") -> bool:
        """
        导出测试结果到Excel
        
        Args:
            test_results: 测试结果列表
            output_path: 输出文件路径
            company_name: 公司名称
            
        Returns:
            是否成功
        """
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
            
            # 创建工作簿
            if self.template_path and os.path.exists(self.template_path):
                wb = openpyxl.load_workbook(self.template_path)
                ws = wb.active
            else:
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "测试报告"
                
                # 设置标题
                self._setup_excel_template(ws, company_name)
                
            # 写入数据
            start_row = 5  # 数据开始行
            for i, result in enumerate(test_results):
                row = start_row + i
                
                # 格式化时间
                dt = datetime.fromtimestamp(result["timestamp"])
                time_str = dt.strftime("%Y-%m-%d %H:%M:%S")
                
                # 写入数据
                ws.cell(row=row, column=1, value=time_str)
                ws.cell(row=row, column=2, value=result["fault_code"])
                ws.cell(row=row, column=3, value=result["fault_name"])
                ws.cell(row=row, column=4, value=result["test_type"])
                ws.cell(row=row, column=5, value=result["status"])
                ws.cell(row=row, column=6, value=result["bms_address"])
                ws.cell(row=row, column=7, value=result["domain"])
                ws.cell(row=row, column=8, value=f"{result['duration_sec']:.2f}")
                ws.cell(row=row, column=9, value=result["error_message"] or "")
                
                # 设置状态颜色
                status_cell = ws.cell(row=row, column=5)
                if result["status"] == "PASS":
                    status_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                else:
                    status_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                    
            # 添加统计信息
            self._add_statistics(ws, test_results, start_row + len(test_results) + 2)
            
            # 保存文件
            wb.save(output_path)
            logger.info(f"Excel报告已导出到: {output_path}")
            return True
            
        except ImportError:
            logger.error("未安装openpyxl库，无法导出Excel")
            return False
        except Exception as e:
            logger.error(f"导出Excel失败: {e}")
            return False
            
    def _setup_excel_template(self, ws, company_name: str):
        """设置Excel模板"""
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        
        # 设置列宽
        column_widths = [20, 15, 40, 10, 10, 10, 10, 15, 40]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
            
        # 设置标题
        title_font = Font(name="微软雅黑", size=16, bold=True)
        ws.merge_cells('A1:I1')
        title_cell = ws.cell(row=1, column=1, value=f"{company_name} - BMS报警测试报告")
        title_cell.font = title_font
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # 设置副标题
        subtitle_font = Font(name="微软雅黑", size=12)
        ws.merge_cells('A2:I2')
        subtitle_cell = ws.cell(row=2, column=1, 
                               value=f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        subtitle_cell.font = subtitle_font
        subtitle_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # 设置表头
        headers = ["测试时间", "故障码", "故障名称", "测试类型", "状态", 
                  "BMS地址", "域类型", "持续时间(秒)", "错误信息"]
        
        header_font = Font(name="微软雅黑", size=11, bold=True)
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font_color = Font(color="FFFFFF", bold=True)
        
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=header)
            cell.font = header_font_color
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            
    def _add_statistics(self, ws, test_results: List[Dict[str, Any]], start_row: int):
        """添加统计信息"""
        from openpyxl.styles import Font, Alignment
        
        if not test_results:
            return
            
        # 计算统计信息
        total = len(test_results)
        passed = sum(1 for r in test_results if r["status"] == "PASS")
        failed = total - passed
        pass_rate = passed / total * 100 if total > 0 else 0
        
        # 添加统计标题
        stat_title_font = Font(name="微软雅黑", size=12, bold=True)
        ws.merge_cells(f'A{start_row}:I{start_row}')
        title_cell = ws.cell(row=start_row, column=1, value="测试统计信息")
        title_cell.font = stat_title_font
        
        # 添加统计内容
        stats = [
            ("总测试数", str(total)),
            ("通过数", str(passed)),
            ("失败数", str(failed)),
            ("通过率", f"{pass_rate:.1f}%")
        ]
        
        for i, (label, value) in enumerate(stats):
            row = start_row + 1 + i
            ws.cell(row=row, column=1, value=label)
            ws.cell(row=row, column=2, value=value)
            
    def export_summary_report(self, statistics: Dict[str, Any], 
                            output_path: str) -> bool:
        """
        导出汇总报告
        
        Args:
            statistics: 统计信息
            output_path: 输出文件路径
            
        Returns:
            是否成功
        """
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill
            
            # 创建工作簿
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "测试汇总"
            
            # 设置标题
            title_font = Font(name="微软雅黑", size=14, bold=True)
            ws.merge_cells('A1:D1')
            title_cell = ws.cell(row=1, column=1, value="BMS报警测试汇总报告")
            title_cell.font = title_font
            title_cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # 设置生成时间
            time_font = Font(name="微软雅黑", size=10)
            ws.merge_cells('A2:D2')
            time_cell = ws.cell(row=2, column=1, 
                               value=f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
            time_cell.font = time_font
            time_cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # 添加总体统计
            ws.cell(row=4, column=1, value="总体统计").font = Font(bold=True)
            
            overall_stats = [
                ("总测试数", statistics.get("total", 0)),
                ("通过数", statistics.get("passed", 0)),
                ("失败数", statistics.get("failed", 0)),
                ("通过率", f"{statistics.get('pass_rate', 0):.1f}%"),
                ("平均耗时", f"{statistics.get('avg_duration', 0):.2f}秒")
            ]
            
            for i, (label, value) in enumerate(overall_stats):
                row = 5 + i
                ws.cell(row=row, column=1, value=label)
                ws.cell(row=row, column=2, value=value)
                
            # 添加故障统计
            if "fault_statistics" in statistics:
                start_row = 5 + len(overall_stats) + 2
                ws.cell(row=start_row, column=1, value="故障统计（Top 10）").font = Font(bold=True)
                
                headers = ["故障码", "故障名称", "测试次数", "通过数", "通过率"]
                for col, header in enumerate(headers, 1):
                    ws.cell(row=start_row + 1, column=col, value=header).font = Font(bold=True)
                    
                for i, fault_stat in enumerate(statistics["fault_statistics"]):
                    row = start_row + 2 + i
                    ws.cell(row=row, column=1, value=fault_stat["fault_code"])
                    ws.cell(row=row, column=2, value=fault_stat["fault_name"])
                    ws.cell(row=row, column=3, value=fault_stat["count"])
                    ws.cell(row=row, column=4, value=fault_stat["passed"])
                    ws.cell(row=row, column=5, value=f"{fault_stat['pass_rate']:.1f}%")
                    
            # 添加每日统计
            if "daily_statistics" in statistics:
                start_row = start_row + len(statistics["fault_statistics"]) + 4
                ws.cell(row=start_row, column=1, value="每日统计（最近7天）").font = Font(bold=True)
                
                headers = ["日期", "测试次数", "通过数", "通过率"]
                for col, header in enumerate(headers, 1):
                    ws.cell(row=start_row + 1, column=col, value=header).font = Font(bold=True)
                    
                for i, daily_stat in enumerate(statistics["daily_statistics"]):
                    row = start_row + 2 + i
                    ws.cell(row=row, column=1, value=daily_stat["date"])
                    ws.cell(row=row, column=2, value=daily_stat["count"])
                    ws.cell(row=row, column=3, value=daily_stat["passed"])
                    ws.cell(row=row, column=4, value=f"{daily_stat['pass_rate']:.1f}%")
                    
            # 调整列宽
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width
                
            # 保存文件
            wb.save(output_path)
            logger.info(f"汇总报告已导出到: {output_path}")
            return True
            
        except ImportError:
            logger.error("未安装openpyxl库，无法导出Excel")
            return False
        except Exception as e:
            logger.error(f"导出汇总报告失败: {e}")
            return False
            
    def export_to_html(self, test_results: List[Dict[str, Any]], 
                      output_path: str) -> bool:
        """
        导出测试结果到HTML
        
        Args:
            test_results: 测试结果列表
            output_path: 输出文件路径
            
        Returns:
            是否成功
        """
        try:
            # 计算统计信息
            total = len(test_results)
            passed = sum(1 for r in test_results if r["status"] == "PASS")
            failed = total - passed
            pass_rate = passed / total * 100 if total > 0 else 0
            
            # 生成HTML
            html_content = f"""
            <!DOCTYPE html>
            <html>
            <head>
                <meta charset="UTF-8">
                <title>BMS报警测试报告</title>
                <style>
                    body {{ font-family: Arial, sans-serif; margin: 20px; }}
                    h1 {{ color: #333; }}
                    .summary {{ background-color: #f5f5f5; padding: 15px; border-radius: 5px; margin-bottom: 20px; }}
                    .summary-item {{ margin: 5px 0; }}
                    .summary-label {{ font-weight: bold; }}
                    table {{ width: 100%; border-collapse: collapse; margin-top: 20px; }}
                    th, td {{ border: 1px solid #ddd; padding: 8px; text-align: left; }}
                    th {{ background-color: #4CAF50; color: white; }}
                    tr:nth-child(even) {{ background-color: #f2f2f2; }}
                    .status-pass {{ color: green; font-weight: bold; }}
                    .status-fail {{ color: red; font-weight: bold; }}
                    .timestamp {{ font-size: 0.9em; color: #666; }}
                </style>
            </head>
            <body>
                <h1>BMS报警测试报告</h1>
                <div class="timestamp">生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</div>
                
                <div class="summary">
                    <h2>测试统计</h2>
                    <div class="summary-item">
                        <span class="summary-label">总测试数:</span> {total}
                    </div>
                    <div class="summary-item">
                        <span class="summary-label">通过数:</span> {passed}
                    </div>
                    <div class="summary-item">
                        <span class="summary-label">失败数:</span> {failed}
                    </div>
                    <div class="summary-item">
                        <span class="summary-label">通过率:</span> {pass_rate:.1f}%
                    </div>
                </div>
                
                <h2>详细测试结果</h2>
                <table>
                    <tr>
                        <th>测试时间</th>
                        <th>故障码</th>
                        <th>故障名称</th>
                        <th>测试类型</th>
                        <th>状态</th>
                        <th>BMS地址</th>
                        <th>域类型</th>
                        <th>持续时间(秒)</th>
                        <th>错误信息</th>
                    </tr>
            """
            
            # 添加数据行
            for result in test_results:
                dt = datetime.fromtimestamp(result["timestamp"])
                time_str = dt.strftime("%Y-%m-%d %H:%M:%S")
                
                status_class = "status-pass" if result["status"] == "PASS" else "status-fail"
                
                html_content += f"""
                    <tr>
                        <td>{time_str}</td>
                        <td>{result["fault_code"]}</td>
                        <td>{result["fault_name"]}</td>
                        <td>{result["test_type"]}</td>
                        <td class="{status_class}">{result["status"]}</td>
                        <td>{result["bms_address"]}</td>
                        <td>{result["domain"]}</td>
                        <td>{result["duration_sec"]:.2f}</td>
                        <td>{result["error_message"] or ""}</td>
                    </tr>
                """
                
            html_content += """
                </table>
            </body>
            </html>
            """
            
            # 写入文件
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(html_content)
                
            logger.info(f"HTML报告已导出到: {output_path}")
            return True
            
        except Exception as e:
            logger.error(f"导出HTML失败: {e}")
            return False