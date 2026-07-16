"""
Excel协议解析器
从Excel信号清单自动生成协议配置
"""

import os
import re
import openpyxl


def _parse_enum_from_desc(desc):
    if not desc:
        return None
    enum_map = {}
    pattern = r'(\d+)\s*[:：]\s*([^,，;；\d]+)'
    matches = re.findall(pattern, desc)
    for val_str, label in matches:
        val = int(val_str)
        label = label.strip().rstrip('；;，,。.')
        if label:
            enum_map[val] = label
    return enum_map if enum_map else None


def _infer_datatype(bit_length, bit_start, factor, description):
    if bit_length <= 8 and bit_start is not None and bit_start > 0:
        return "bitfield"
    if bit_length <= 8:
        if description and ('0:' in description or '0：' in description):
            return "bitfield"
        return "uint8"
    if bit_length == 16:
        if description and ('负值' in description or '正值为充电' in description):
            return "int16"
        return "uint16"
    if bit_length < 16:
        return "bitfield"
    return "uint16"


def parse_excel(filepath):
    if not os.path.exists(filepath):
        return None

    wb = openpyxl.load_workbook(filepath, data_only=True)
    filename = os.path.basename(filepath)

    ws = None
    for sheet in wb.worksheets:
        if sheet.max_column >= 10 and sheet.max_row > 10:
            ws = sheet
            break
    if ws is None:
        ws = wb.worksheets[-1]

    REG_UINT16 = "uint16"
    REG_INT16 = "int16"
    REG_UINT8 = "uint8"
    REG_BITFIELD = "bitfield"

    RUN_STATE_ENUM = {
        0: "无效", 1: "自检", 2: "放电", 3: "充电", 4: "充电完成",
        5: "急停", 6: "系统保护", 7: "充电保护", 8: "放电保护",
        9: "静置", 10: "关机", 11: "待机", 12: "预充", 13: "充电加热",
        14: "簇上电中", 15: "簇压均衡中"
    }
    ALARM_LEVEL_ENUM = {0: "无报警", 1: "一级报警", 2: "二级报警", 3: "三级报警"}
    INSULATION_ENUM = {0: "正/负对地相等", 1: "正对地绝缘电阻小", 2: "负对地绝缘电阻小"}
    RELAY_STATE_ENUM = {0: "断开", 1: "闭合"}
    CHRG_GUN_ENUM = {0: "未插入", 1: "已插入"}
    RELAY_FAULT_ENUM = {0: "无报警", 1: "开路", 2: "粘连", 3: "预留"}

    SPECIAL_ENUMS = {
        "系统运行状态": RUN_STATE_ENUM,
        "组当前最大报警等级": ALARM_LEVEL_ENUM,
        "绝缘电阻极性": INSULATION_ENUM,
        "组当前故障": None,
    }

    signals = []
    control_signals = []
    current_addr = None
    current_category = None
    is_control_section = False

    for row in ws.iter_rows(min_row=2, max_row=300, values_only=True):
        sig_name = row[2]
        if sig_name is None:
            continue
        sig_name = str(sig_name).strip()

        cat = row[0]
        if cat:
            current_category = str(cat).strip()
            if '控制指令' in current_category:
                is_control_section = True
            elif current_category and '控制指令' not in current_category:
                is_control_section = False

        addr_val = row[3]
        bit_start = int(row[5]) if row[5] is not None else 0

        if isinstance(addr_val, (int, float)):
            current_addr = int(addr_val)
        elif isinstance(addr_val, str) and addr_val.startswith('='):
            if current_addr is not None:
                current_addr += 1
            else:
                current_addr = 1
        elif addr_val is None:
            if current_addr is None:
                current_addr = 1
            elif bit_start > 0:
                pass
            else:
                current_addr += 1

        bit_length = int(row[4]) if row[4] else 16
        bit_start = int(row[5]) if row[5] is not None else 0
        description = str(row[7]) if row[7] else ''
        factor = row[8] if row[8] is not None else 1
        offset = row[9] if row[9] is not None else 0
        unit = row[10] if row[10] is not None else ''
        en_name = str(row[12]) if row[12] else ''

        if isinstance(factor, str):
            try:
                factor = float(factor)
            except ValueError:
                factor = 1
        if isinstance(offset, str):
            try:
                offset = float(offset)
            except ValueError:
                offset = 0

        if isinstance(unit, (int, float)) and unit != 0:
            next_val = row[11] if row[11] is not None else ''
            next_val_str = str(next_val).strip()
            if next_val_str and not next_val_str.replace('.', '').replace('-', '').isdigit():
                offset = unit
                unit = next_val_str
            else:
                unit = str(unit) if unit != '' else ''
        else:
            unit = str(unit) if unit != '' else ''

        if unit == 'KΩ' or unit == 'KΩ':
            unit = 'kΩ'

        if is_control_section:
            enum_map = _parse_enum_from_desc(description)
            control_signals.append({
                'addr': current_addr,
                'name': sig_name,
                'bit_start': bit_start,
                'bit_length': bit_length,
                'description': description,
                'enum': enum_map,
                'en_name': en_name,
            })
            continue

        if current_addr < 20:
            continue

        datatype = _infer_datatype(bit_length, bit_start, factor, description)

        enum_map = SPECIAL_ENUMS.get(sig_name)
        if enum_map is None and sig_name not in SPECIAL_ENUMS:
            parsed_enum = _parse_enum_from_desc(description)
            if parsed_enum and len(parsed_enum) >= 2:
                enum_map = parsed_enum
            elif '继电器' in sig_name and ('故障' in sig_name or '报警' in sig_name):
                enum_map = RELAY_FAULT_ENUM
            elif '告警' in sig_name or '报警' in sig_name or '过放' in sig_name or '过充' in sig_name or '过温' in sig_name:
                if bit_length <= 2:
                    enum_map = ALARM_LEVEL_ENUM
            elif '继电器状态' in sig_name or '插枪状态' in sig_name:
                enum_map = RELAY_STATE_ENUM

        signals.append({
            'category': current_category,
            'addr': current_addr,
            'name': sig_name,
            'datatype': datatype,
            'bit_start': bit_start,
            'bit_length': bit_length,
            'factor': factor,
            'offset': offset,
            'unit': unit,
            'description': description,
            'enum': enum_map,
            'en_name': en_name,
        })

    wb.close()

    name_count = {}
    for sig in signals:
        name = sig['name']
        if name in name_count:
            name_count[name] += 1
            sig['name'] = f"{name}_{name_count[name]}"
        else:
            name_count[name] = 1

    for sig in signals:
        if sig['category'] == '单体温度':
            addr = sig['addr']
            if 114 <= addr <= 116:
                sig['category'] = '告警状态'
            elif 117 <= addr <= 120:
                sig['category'] = '故障状态'

    for sig in signals:
        name = sig['name']
        cat = sig['category']
        if name.startswith('预留'):
            sig['category'] = '预留'
            continue
        if cat in ('单体电压', '单体温度'):
            if any(kw in name for kw in ['最大单体电压值', '最大单体电压箱号', '最大单体电压编号',
                                          '最小单体电压值', '最小单体电压箱号', '最小单体电压编号',
                                          '最大单体温度值', '最大单体温度箱号', '最大单体温度编号',
                                          '最小单体温度值', '最小单体温度箱号', '最小单体温度编号']):
                sig['category'] = '极值'
                continue
        if cat in ('告警状态', '故障状态', '故障信息'):
            sig['category'] = '告警'

    signals.sort(key=lambda s: (s['category'] or '', s['addr'], s['bit_start']))

    categories = []
    current_cat = None
    for sig in signals:
        cat = sig['category']
        if cat is None:
            cat = current_cat if current_cat else "未分类"
        if cat != current_cat:
            current_cat = cat
            categories.append((cat, []))
        categories[-1][1].append(sig['name'])

    tab_order = ['请求指令信息', '组总体状态', '电池组当前状态', '电流情况', '电压情况', '极值', '告警', '预留']
    tab_dict = {cat: names for cat, names in categories if cat and '控制指令' not in cat}

    status_tabs = []
    for tab in tab_order:
        if tab in tab_dict:
            status_tabs.append((tab, tab_dict[tab]))
    for cat, names in tab_dict.items():
        if cat not in tab_order:
            status_tabs.append((cat, names))

    register_groups = {}
    current_cat = None
    for sig in signals:
        cat = sig['category']
        if cat is None:
            cat = current_cat if current_cat else "未分类"
        current_cat = cat
        if cat not in register_groups:
            register_groups[cat] = []
        register_groups[cat].append((
            sig['addr'], sig['name'], sig['datatype'],
            sig['bit_start'], sig['bit_length'],
            sig['factor'], sig['offset'], sig['unit'],
            sig['description'], sig['enum'], sig['en_name']
        ))

    register_groups = {k: v for k, v in register_groups.items()
                       if k and '控制指令' not in k}

    all_addrs = sorted(set(sig['addr'] for sig in signals))
    read_addrs = [a for a in all_addrs if a >= 20]
    poll_groups = []
    if read_addrs:
        start = read_addrs[0]
        end = start
        for addr in read_addrs[1:]:
            if addr - start < 20 and addr == end + 1:
                end = addr
            else:
                poll_groups.append((start, end - start + 1))
                start = addr
                end = addr
        poll_groups.append((start, end - start + 1))

    return {
        'register_groups': register_groups,
        'poll_groups': poll_groups,
        'status_tabs': status_tabs,
        'control_signals': control_signals,
        'enums': {
            'RUN_STATE_ENUM': RUN_STATE_ENUM,
            'ALARM_LEVEL_ENUM': ALARM_LEVEL_ENUM,
            'INSULATION_ENUM': INSULATION_ENUM,
            'RELAY_STATE_ENUM': RELAY_STATE_ENUM,
            'CHRG_GUN_ENUM': CHRG_GUN_ENUM,
            'RELAY_FAULT_ENUM': RELAY_FAULT_ENUM,
        },
        'source_info': f'Excel: {filename}',
    }
